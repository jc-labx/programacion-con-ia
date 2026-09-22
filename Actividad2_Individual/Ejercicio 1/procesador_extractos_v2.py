#! python
#
# Programación con Inteligencia Artificial - Actividad individual - Unidad 2
#
# Autor: Jaime Alberto Chica Betancourt
#
# Universidad de Investigación y Desarrollo – UDI
# Maestría en Ciencia de Datos e Inteligencia Artificial
# 7613 – Programación con IA
# Docente: Julián Andrés Ramirez Bautista
#
# Fecha de Entrega: 27 de Septiembre de 2026
#
# Consolidación de extractos bancarios mediante POO - v2
#
# Mejoras respecto a la v1:
# - Nueva clase ClasificadorMovimientos.
# - Aplicación del principio de responsabilidad única (SRP).
# - Las reglas de negocio quedan desacopladas del proceso de transformación.
# - Mayor facilidad para agregar nuevas categorías.
#

from __future__ import annotations

from abc import ABC, abstractmethod
from pathlib import Path

import pandas as pd

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


COLUMNAS_NORMALIZADAS = [
    "Fecha",
    "Descripcion",
    "Importe",
    "Moneda",
]

COLUMNAS_SALIDA = [
    "Fecha",
    "Tipo de Gasto",
    "Detalle",
    "Valor",
    "Moneda",
    "Comentario",
]


# ============================================================
# LECTORES
# ============================================================

class LectorExtracto(ABC):
    """Contrato común para todos los bancos."""

    moneda: str

    @abstractmethod
    def leer(self, ruta: str | Path) -> pd.DataFrame:
        raise NotImplementedError

    @staticmethod
    def _comprobar_archivo(ruta: str | Path) -> Path:
        archivo = Path(ruta)

        if not archivo.exists():
            raise FileNotFoundError(
                f"No se encontró el archivo: {archivo}"
            )

        return archivo

    @staticmethod
    def _validar_columnas(
        datos: pd.DataFrame,
        requeridas: set[str]
    ) -> None:

        faltantes = requeridas - set(datos.columns)

        if faltantes:
            raise ValueError(
                f"Faltan columnas obligatorias: {sorted(faltantes)}"
            )

    @staticmethod
    def _importe_texto(
        serie: pd.Series,
        decimal: str,
        miles: str | None = None
    ) -> pd.Series:

        texto = serie.astype(str).str.strip()

        if miles:
            texto = texto.str.replace(
                miles,
                "",
                regex=False
            )

        if decimal != ".":
            texto = texto.str.replace(
                decimal,
                ".",
                regex=False
            )

        return pd.to_numeric(
            texto,
            errors="raise"
        )

    @staticmethod
    def _normalizar(
        fecha: pd.Series,
        descripcion: pd.Series,
        importe: pd.Series,
        moneda: str
    ) -> pd.DataFrame:

        return pd.DataFrame(
            {
                "Fecha": fecha,
                "Descripcion": descripcion.fillna("")
                .astype(str)
                .str.strip(),
                "Importe": importe,
                "Moneda": moneda,
            }
        )[COLUMNAS_NORMALIZADAS]


class LectorBancoA(LectorExtracto):

    moneda = "COP"

    def leer(self, ruta: str | Path) -> pd.DataFrame:

        archivo = self._comprobar_archivo(ruta)

        datos = pd.read_csv(
            archivo,
            sep=";",
            dtype=str,
            encoding="utf-8-sig"
        )

        self._validar_columnas(
            datos,
            {"Fecha", "Descripción", "Importe"}
        )

        fecha = pd.to_datetime(
            datos["Fecha"],
            format="%d/%m/%Y",
            errors="raise"
        )

        importe = self._importe_texto(
            datos["Importe"],
            decimal=",",
            miles="."
        )

        return self._normalizar(
            fecha,
            datos["Descripción"],
            importe,
            self.moneda
        )


class LectorBancoB(LectorExtracto):

    moneda = "EUR"

    def leer(self, ruta: str | Path) -> pd.DataFrame:

        archivo = self._comprobar_archivo(ruta)

        datos = pd.read_csv(
            archivo,
            sep=";",
            dtype=str,
            encoding="utf-8-sig"
        )

        self._validar_columnas(
            datos,
            {"Buchungstag", "Buchungstext", "Betrag"}
        )

        fecha = pd.to_datetime(
            datos["Buchungstag"],
            format="%d.%m.%Y",
            errors="raise"
        )

        importe = self._importe_texto(
            datos["Betrag"],
            decimal=".",
            miles=","
        )

        return self._normalizar(
            fecha,
            datos["Buchungstext"],
            importe,
            self.moneda
        )


class LectorBancoC(LectorExtracto):

    moneda = "USD"

    def leer(self, ruta: str | Path) -> pd.DataFrame:

        archivo = self._comprobar_archivo(ruta)

        datos = pd.read_csv(
            archivo,
            sep=",",
            dtype=str,
            encoding="utf-8-sig",
            quotechar='"'
        )

        self._validar_columnas(
            datos,
            {"Date", "Description", "Amount"}
        )

        fecha = pd.to_datetime(
            datos["Date"],
            format="%m/%d/%Y",
            errors="raise"
        )

        importe = self._importe_texto(
            datos["Amount"],
            decimal=",",
            miles="."
        )

        return self._normalizar(
            fecha,
            datos["Description"],
            importe,
            self.moneda
        )


class LectorBancoD(LectorExtracto):

    moneda = "COP"

    def leer(self, ruta: str | Path) -> pd.DataFrame:

        archivo = self._comprobar_archivo(ruta)

        datos = pd.read_csv(
            archivo,
            sep=";",
            dtype=str,
            encoding="utf-8-sig"
        )

        self._validar_columnas(
            datos,
            {"Fecha", "Transaccion", "Varlor"}
        )

        fecha = pd.to_datetime(
            datos["Fecha"],
            format="%d-%m-%Y",
            errors="raise"
        )

        importe = self._importe_texto(
            datos["Varlor"],
            decimal=",",
            miles="."
        )

        return self._normalizar(
            fecha,
            datos["Transaccion"],
            importe,
            self.moneda
        )


# ============================================================
# CLASIFICADOR (NUEVO EN V2)
# ============================================================

class ClasificadorMovimientos:
    """
    Encapsula las reglas de clasificación.
    """

    REGLAS = {
        "EXITO": ("Alimentación", "Éxito"),
        "CARULLA": ("Alimentación", "Carulla"),
        "LIDL": ("Alimentación", "Lidl"),
        "REWE": ("Alimentación", "Rewe"),
        "WERKSTATT": ("Transporte", "Taller"),
        "UBER": ("Transporte", "Uber"),
        "SHELL": ("Transporte", "Combustible"),
        "NETFLIX": ("Servicios", "Netflix"),
        "VODAFONE": ("Servicios", "Vodafone"),
        "AMAZON": ("Compras", "Amazon"),
        "FARMACIA": ("Salud", "Farmacia"),
        "APOTHEKE": ("Salud", "Farmacia"),
        "RESTAURANT": ("Restaurantes", "Restaurante"),
    }

    def clasificar(
        self,
        descripcion: str
    ) -> tuple[str, str\]:

        texto = str(descripcion).upper()

        for palabra_clave, resultado in self.REGLAS.items():

            if palabra_clave in texto:
                return resultado

        return (
            "Revisar",
            descripcion[:80]
        )


# ============================================================
# TRANSFORMADOR
# ============================================================

class TransformadorMovimientos:

    def __init__(
        self,
        clasificador: ClasificadorMovimientos
    ) -> None:

        self.clasificador = clasificador

    def transformar(
        self,
        datos: pd.DataFrame
    ) -> pd.DataFrame:

        tipos_gasto = []
        detalles = []

        for descripcion in datos["Descripcion"\]:

            categoria, detalle = (
                self.clasificador.clasificar(
                    descripcion
                )
            )

            tipos_gasto.append(categoria)
            detalles.append(detalle)

        return pd.DataFrame(
            {
                "Fecha": datos["Fecha"],
                "Tipo de Gasto": tipos_gasto,
                "Detalle": detalles,
                "Valor": datos["Importe"],
                "Moneda": datos["Moneda"],
                "Comentario": "",
            }
        )


# ============================================================
# ESCRITOR EXCEL
# ============================================================

class EscritorExcel:

    def escribir(
        self,
        datos: pd.DataFrame,
        ruta: str | Path
    ) -> None:

        archivo = Path(ruta)

        if archivo.suffix.lower() != ".xlsx":
            raise ValueError(
                "La salida debe tener extensión .xlsx"
            )

        with pd.ExcelWriter(
            archivo,
            engine="openpyxl",
            date_format="DD.MM.YYYY"
        ) as writer:

            datos.to_excel(
                writer,
                sheet_name="Movimientos",
                index=False
            )

        libro = load_workbook(archivo)
        hoja = libro["Movimientos"]

        hoja.freeze_panes = "A2"
        hoja.sheet_view.showGridLines = False

        relleno = PatternFill(
            "solid",
            fgColor="1F4E78"
        )

        for celda in hoja[1\]:
            celda.fill = relleno
            celda.font = Font(
                color="FFFFFF",
                bold=True
            )
            celda.alignment = Alignment(
                horizontal="center"
            )

        anchos = {
            "A": 14,
            "B": 20,
            "C": 24,
            "D": 18,
            "E": 12,
            "F": 34,
        }

        for columna, ancho in anchos.items():
            hoja.column_dimensions[columna].width = ancho

        for celda in hoja["A"][1:\]:
            celda.number_format = "DD.MM.YYYY       for celda in hoja["D"][1:]:
            celda.number_format = (
                '#,##0.00;[Red](#,##0.00);-'
            )

        tabla = Table(
            displayName="TablaMovimientos",
            ref=hoja.dimensions
        )

        tabla.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )

        hoja.add_table(tabla)

        hoja.auto_filter.ref = None

        libro.save(archivo)


# ============================================================
# PROCESADOR
# ============================================================

class ProcesadorExtractos:

    def __init__(
        self,
        transformador: TransformadorMovimientos
    ) -> None:

        self.transformador = transformador

    def procesar(
        self,
        lector: LectorExtracto,
        ruta: str | Path
    ) -> pd.DataFrame:

        datos = lector.leer(ruta)

        return self.transformador.transformar(
            datos
        )


# ============================================================
# MAIN
# ============================================================

if __name__ == "__main__":

    base = Path(__file__).resolve().parent

    clasificador = ClasificadorMovimientos()

    transformador = TransformadorMovimientos(
        clasificador
    )

    procesador = ProcesadorExtractos(
        transformador
    )

    entradas = [
        (LectorBancoA(), base / "extracto_banco_a.csv"),
        (LectorBancoB(), base / "extracto_banco_b.csv"),
        (LectorBancoC(), base / "extracto_banco_c.csv"),
        (LectorBancoD(), base / "extracto_banco_d.csv"),
    ]

    resultados = [
        procesador.procesar(
            lector,
            ruta
        )
        for lector, ruta in entradas
    ]

    consolidado = pd.concat(
        resultados,
        ignore_index=True
    ).sort_values("Fecha")

    EscritorExcel().escribir(
        consolidado,
        base / "resultados_consolidados.xlsx"
    )

    print(
        f"Movimientos procesados: {len(consolidado)}"
    )

    print(
        f"Movimientos para revisar: "
        f"{(consolidado['Tipo de Gasto'] == 'Revisar').sum()}"
    )